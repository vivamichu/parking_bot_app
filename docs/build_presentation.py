"""Generate docs/presentation.pptx for Stage 2.

Run from the stage 2 root:  python docs/build_presentation.py
Rebuilds the deck from the screenshots in docs/ and the explanatory content below.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

DOCS = Path(__file__).resolve().parent

# 16:9 widescreen
EMU_W, EMU_H = Inches(13.333), Inches(7.5)

# Palette
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
BLUE = RGBColor(0x1E, 0x5A, 0x96)
ACCENT = RGBColor(0x2E, 0x86, 0xDE)
LIGHT = RGBColor(0xF4, 0xF7, 0xFB)
GREY = RGBColor(0x5A, 0x66, 0x75)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x27, 0xAE, 0x60)

SHOTS = {
    "book": DOCS / "Screenshot 2026-07-03 at 14.01.07.png",
    "admin": DOCS / "Screenshot 2026-07-03 at 14.01.30.png",
    "confirm": DOCS / "Screenshot 2026-07-03 at 14.01.47.png",
}

prs = Presentation()
prs.slide_width, prs.slide_height = EMU_W, EMU_H
BLANK = prs.slide_layouts[6]


def _bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _box(slide, x, y, w, h):
    return slide.shapes.add_textbox(x, y, w, h).text_frame


def _run(p, text, size, color, bold=False, italic=False):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = "Calibri"
    return r


def _band(slide, color, y, h):
    from pptx.enum.shapes import MSO_SHAPE

    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, y, EMU_W, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def title_slide():
    s = prs.slides.add_slide(BLANK)
    _bg(s, NAVY)
    _band(s, BLUE, Inches(2.55), Inches(0.06))
    tf = _box(s, Inches(0.9), Inches(2.5), Inches(11.5), Inches(2.6))
    p = tf.paragraphs[0]
    _run(p, "🅿️  Astana Central Parking", 40, WHITE, bold=True)
    p2 = tf.add_paragraph()
    _run(p2, "Stage 2 — Human-in-the-Loop Admin Approval", 26, ACCENT, bold=True)
    p2.space_before = Pt(10)
    p3 = tf.add_paragraph()
    _run(
        p3,
        "A second LangChain agent that escalates every reservation to a human "
        "administrator and returns the decision to the visitor.",
        16,
        LIGHT,
    )
    p3.space_before = Pt(16)
    foot = _box(s, Inches(0.9), Inches(6.5), Inches(11.5), Inches(0.6))
    _run(
        foot.paragraphs[0],
        "LangChain · LangGraph · FastAPI · Streamlit · Milvus Lite · SQLite",
        13,
        GREY,
    )


def content_slide(title, kicker, bullets, note=None):
    s = prs.slides.add_slide(BLANK)
    _bg(s, WHITE)
    _band(s, NAVY, 0, Inches(1.35))
    ktf = _box(s, Inches(0.7), Inches(0.28), Inches(12), Inches(0.4))
    _run(ktf.paragraphs[0], kicker.upper(), 12, ACCENT, bold=True)
    ttf = _box(s, Inches(0.7), Inches(0.62), Inches(12), Inches(0.7))
    _run(ttf.paragraphs[0], title, 28, WHITE, bold=True)

    body = _box(s, Inches(0.85), Inches(1.75), Inches(11.6), Inches(5.2))
    body.word_wrap = True
    first = True
    for text, level in bullets:
        p = body.paragraphs[0] if first else body.add_paragraph()
        first = False
        p.level = level
        bullet = "▸ " if level == 0 else "• "
        _run(p, bullet, 16, ACCENT if level == 0 else BLUE, bold=True)
        _run(p, text, 17 if level == 0 else 15, NAVY if level == 0 else GREY,
             bold=(level == 0))
        p.space_after = Pt(7 if level == 0 else 3)
    if note:
        ntf = _box(s, Inches(0.85), Inches(6.75), Inches(11.6), Inches(0.5))
        _run(ntf.paragraphs[0], note, 13, GREY, italic=True)


def screenshot_slide(title, kicker, img_path, caption):
    s = prs.slides.add_slide(BLANK)
    _bg(s, LIGHT)
    _band(s, NAVY, 0, Inches(1.15))
    ktf = _box(s, Inches(0.7), Inches(0.2), Inches(12), Inches(0.35))
    _run(ktf.paragraphs[0], kicker.upper(), 12, ACCENT, bold=True)
    ttf = _box(s, Inches(0.7), Inches(0.5), Inches(12), Inches(0.6))
    _run(ttf.paragraphs[0], title, 25, WHITE, bold=True)

    # Fit image: source is 2880x1800 (16:10). Available area below header.
    avail_top = Inches(1.45)
    avail_h = Inches(5.35)
    src_w, src_h = 2880, 1800
    max_w = Inches(11.8)
    # scale to height first
    w = Emu(int(avail_h * src_w / src_h))
    h = avail_h
    if w > max_w:
        w = max_w
        h = Emu(int(max_w * src_h / src_w))
    x = Emu(int((EMU_W - w) / 2))
    y = Emu(int(avail_top + (avail_h - h) / 2))
    if img_path.exists():
        s.shapes.add_picture(str(img_path), x, y, width=w, height=h)
    else:
        _run(_box(s, x, y, w, Inches(1)).paragraphs[0],
             f"[missing image: {img_path.name}]", 14, GREY)

    ctf = _box(s, Inches(0.7), Inches(6.85), Inches(12), Inches(0.5))
    cp = ctf.paragraphs[0]
    cp.alignment = PP_ALIGN.CENTER
    _run(cp, caption, 14, BLUE, bold=True)


def flow_slide():
    s = prs.slides.add_slide(BLANK)
    _bg(s, WHITE)
    _band(s, NAVY, 0, Inches(1.35))
    _run(_box(s, Inches(0.7), Inches(0.28), Inches(12), Inches(0.4)).paragraphs[0],
         "END-TO-END", 12, ACCENT, bold=True)
    _run(_box(s, Inches(0.7), Inches(0.62), Inches(12), Inches(0.7)).paragraphs[0],
         "The reservation lifecycle", 28, WHITE, bold=True)

    from pptx.enum.shapes import MSO_SHAPE

    steps = [
        ("1  COLLECT", "Agent 1 gathers name, plate & period, then validates them.", BLUE),
        ("2  ESCALATE", "A PENDING row is written, tagged with the chat session_id.", BLUE),
        ("3  REVIEW", "Admin opens the web page and clicks Approve / Reject.", ACCENT),
        ("4  DECIDE", "Agent 2 maps the decision → status and updates the record.", ACCENT),
        ("5  NOTIFY", "The outbox is polled; the visitor sees CONFIRMED / REFUSED.", GREEN),
    ]
    n = len(steps)
    gap = Inches(0.25)
    total_w = EMU_W - Inches(1.4)
    card_w = Emu(int((total_w - gap * (n - 1)) / n))
    y = Inches(2.4)
    card_h = Inches(3.0)
    x = Inches(0.7)
    for label, desc, col in steps:
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT
        card.line.color.rgb = col
        card.line.width = Pt(1.5)
        card.shadow.inherit = False
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.15)
        tf.margin_right = Inches(0.15)
        tf.margin_top = Inches(0.2)
        _run(tf.paragraphs[0], label, 15, col, bold=True)
        p = tf.add_paragraph()
        _run(p, desc, 12, GREY)
        p.space_before = Pt(8)
        x = Emu(int(x + card_w + gap))
    _run(_box(s, Inches(0.7), Inches(5.9), Inches(12), Inches(0.5)).paragraphs[0],
         "Agents 1 and 2 never call each other — they cooperate through the shared SQLite store.",
         14, NAVY, italic=True)


# ----------------------------- Build the deck ----------------------------- #
title_slide()

content_slide(
    "The problem & the goal", "Motivation",
    [
        ("Stage 1 could take a booking — but nobody approved it", 0),
        ("Reservations were created as PENDING with no human in the loop", 1),
        ("Stage 2 goal: connect a human administrator for approval", 0),
        ("Send a reservation request to an administrator", 1),
        ("Receive an approve / reject decision back", 1),
        ("Keep the visitor's chatbot and the admin in sync — automatically", 1),
        ("Outcome: an automated system with a human decision-maker in the loop", 0),
    ],
)

content_slide(
    "Two agents, one shared store", "Architecture",
    [
        ("Agent 1 — Visitor chatbot (LangGraph ReAct, gpt-4o-mini)", 0),
        ("RAG over the knowledge base + live SQLite tools + make_reservation", 1),
        ("Guardrails: prompt-injection filter in, PII scrub out", 1),
        ("Agent 2 — Admin agent (LangChain ReAct)", 0),
        ("Tools: get_reservation_details, check_space_availability, apply_decision", 1),
        ("Verifies a space is still free, then confirms or refuses", 1),
        ("They communicate through SQLite — not direct calls", 0),
        ("reservations table = request queue (a PENDING row = an escalation)", 1),
        ("notifications table = reply channel (keyed to the visitor's session_id)", 1),
    ],
    note="FastAPI admin server is the human's window into the queue: GET /admin/ · POST /admin/decision",
)

flow_slide()

screenshot_slide(
    "1 · Visitor books a space", "Agent 1 — Streamlit chat",
    SHOTS["book"],
    "The chatbot collects the details and replies: reservation is PENDING administrator confirmation.",
)

screenshot_slide(
    "2 · Administrator approves via REST", "Agent 2 — FastAPI admin page",
    SHOTS["admin"],
    "http://localhost:8600/admin/ lists the pending request with one-click Approve / Reject (uvicorn serving below).",
)

screenshot_slide(
    "3 · Decision returns to the visitor", "Async notification",
    SHOTS["confirm"],
    "The admin agent updates the status and pushes “Reservation #1 CONFIRMED” back into the visitor's chat.",
)

content_slide(
    "Engineering quality", "Testing · CI/CD",
    [
        ("50 automated pytest tests — ≥ 2 per module, fully offline", 0),
        ("New: test_admin_agent (decision tools) & test_admin_server (FastAPI TestClient)", 1),
        ("LLM stubbed; isolated SQLite fixtures keep every test hermetic", 1),
        ("CI/CD — GitHub Actions (.github/workflows/ci.yml)", 0),
        ("Matrix on Python 3.11 & 3.12; installs lean deps + spaCy model; runs pytest", 1),
        ("Runs on every push and pull request", 1),
        ("Robustness fixes surfaced along the way", 0),
        ("Async notification no longer interrupted by the auto-refresh timer", 1),
        ("python-multipart added so the admin form endpoint works in a clean env", 1),
    ],
)

content_slide(
    "Outcome", "Summary",
    [
        ("A working two-agent, human-in-the-loop approval system", 0),
        ("Visitor books → admin approves/rejects → visitor is notified, end to end", 1),
        ("Reservation requests are generated and sent to the administrator", 0),
        ("Responses are received and applied (confirmed / cancelled)", 0),
        ("Communication between the two agents is maintained via the shared DB", 0),
        ("Documented, tested, and continuously integrated", 0),
    ],
    note="Run:  uvicorn parking_bot.admin_server:app --port 8600   +   streamlit run app.py",
)

out = DOCS / "presentation.pptx"
prs.save(str(out))
print("Wrote", out, "with", len(prs.slides._sldIdLst), "slides")
